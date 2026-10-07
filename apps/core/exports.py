"""
Exports Excel (.xlsx) et PDF de tableaux (contrôles qualité, indicateurs...).
Les deux formats partagent la même définition : un titre, des en-têtes et
des lignes de valeurs.
"""

from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _valeur_excel(valeur):
    if isinstance(valeur, Decimal):
        return float(valeur)
    if hasattr(valeur, "tzinfo") and getattr(valeur, "tzinfo", None) is not None:
        return timezone.localtime(valeur).replace(tzinfo=None)
    return valeur


def classeur(feuilles, nom_fichier):
    """
    feuilles = [(titre, entetes, lignes), ...] -> réponse .xlsx (une feuille
    par tableau, en-têtes en gras, filtres automatiques, largeurs ajustées).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .models import ParametreEntreprise
    couleur = ParametreEntreprise.courant().couleur.lstrip("#")
    classeur_xlsx = Workbook()
    classeur_xlsx.remove(classeur_xlsx.active)
    for titre, entetes, lignes in feuilles:
        feuille = classeur_xlsx.create_sheet(titre[:31])
        feuille.append(list(entetes))
        for cellule in feuille[1]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor=couleur)
        for ligne in lignes:
            feuille.append([_valeur_excel(v) for v in ligne])
        for index, entete in enumerate(entetes, start=1):
            largeur = max([len(str(entete))] + [len(str(l[index - 1] or "")) for l in lignes[:500]])
            feuille.column_dimensions[get_column_letter(index)].width = min(max(largeur + 2, 8), 60)
        feuille.freeze_panes = "A2"
        if lignes:
            feuille.auto_filter.ref = feuille.dimensions
    tampon = BytesIO()
    classeur_xlsx.save(tampon)
    reponse = HttpResponse(tampon.getvalue(), content_type=TYPE_XLSX)
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}.xlsx"'
    return reponse


def tableau_pdf(titre, reference, entetes, lignes, largeurs, utilisateur, colonnes_nombres=(), texte_avant=""):
    """Même tableau en PDF (mise en page EVAM : logo, mentions, pagination)."""
    from .pdf import DocumentPDF
    document = DocumentPDF(titre, reference, timezone.now(), utilisateur)
    document.avec_texte(texte_avant)
    document.avec_lignes(entetes, [[("" if v is None else v) for v in ligne] for ligne in lignes], largeurs,
                         colonnes_nombres=colonnes_nombres)
    return document
